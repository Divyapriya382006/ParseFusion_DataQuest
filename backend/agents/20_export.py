from __future__ import annotations

import base64
import csv
import html as _html
import io
import json
import re
import time
import uuid
from datetime import timedelta
from typing import Any, Optional
from urllib.parse import quote

from fastapi import APIRouter, Body, Query, Request
from fastapi.responses import JSONResponse, Response
from pydantic import BaseModel, ConfigDict
from sqlalchemy import Column, Integer, String, Table, Text, insert, select

try:
    from . import e_common as ec
except ImportError:  # pragma: no cover
    import e_common as ec

AGENT = "20-export"
VERSION = "1"

exports_tbl = Table(
    "exports", ec.APP_META,
    Column("export_id", String(36), primary_key=True),
    Column("tenant_id", String(64), nullable=False), Column("user_id", String(128), nullable=False, index=True),
    Column("scope", Text, nullable=False), Column("format", String(16), nullable=False),
    Column("content_hash", String(64), nullable=False), Column("size_bytes", Integer, nullable=False),
    Column("masked", Integer, nullable=False), Column("created_at", String(32), nullable=False),
    Column("content_type", String(100)), Column("filename", String(160)),
)

DEFAULT_SCOPES = ["source", "case", "batch", "finding", "action"]
DEFAULT_FORMATS = ["json", "markdown", "csv", "xlsx", "pdf", "html", "docx"]
DEFAULT_SENSITIVE_KEYS = ["email", "phone", "mobile", "account", "account_number", "iban", "tax_id", "pan", "aadhaar", "ssn", "address", "dob", "password", "token"]
DEFAULT_EVIDENCE_KEYS = ["evidence", "evidence_references", "supporting_evidence"]
_EMAIL = re.compile(r"[\w.+-]+@[\w-]+(\.[\w-]+)+")
_LONGNUM = re.compile(r"\d{9,}")


class ExportOptions(BaseModel):
    model_config = ConfigDict(extra="forbid")
    masked: Optional[bool] = None
    include_evidence: bool = False


class ExportScope(BaseModel):
    model_config = ConfigDict(extra="forbid")
    type: str
    ids: list


class ExportIn(BaseModel):
    model_config = ConfigDict(extra="forbid")
    scope: ExportScope
    format: str
    options: Optional[ExportOptions] = None


# ---------------------------------------------------------------------------
# data preparation
# ---------------------------------------------------------------------------
def _walk(obj: Any, fn_key) -> Any:
    """Rebuilds obj dropping keys for which fn_key(key) is True."""
    if isinstance(obj, dict):
        return {k: _walk(v, fn_key) for k, v in obj.items() if not fn_key(str(k))}
    if isinstance(obj, list):
        return [_walk(v, fn_key) for v in obj]
    return obj


def _mask_value(v: Any) -> Any:
    if isinstance(v, str):
        s = _EMAIL.sub("[email]", v)
        return _LONGNUM.sub(lambda m: "*" * (len(m.group(0)) - 2) + m.group(0)[-2:], s)
    return v


def _mask(obj: Any, sensitive: set) -> Any:
    if isinstance(obj, dict):
        out = {}
        for k, v in obj.items():
            if str(k).lower() in sensitive and not isinstance(v, (dict, list)):
                s = "" if v is None else str(v)
                out[k] = ("*" * max(len(s) - 2, 3) + s[-2:]) if s else v
            else:
                out[k] = _mask(v, sensitive)
        return out
    if isinstance(obj, list):
        return [_mask(v, sensitive) for v in obj]
    return _mask_value(obj)


def _cell(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, (dict, list)):
        return json.dumps(v, sort_keys=True, ensure_ascii=False, separators=(",", ":"))
    return str(v)


def _flatten(prefix: str, obj: Any, kv: list, tables: list) -> None:
    if isinstance(obj, dict):
        for k in obj:
            _flatten(f"{prefix}.{k}" if prefix else str(k), obj[k], kv, tables)
    elif isinstance(obj, list) and obj and all(isinstance(x, dict) for x in obj):
        cols: list = []
        for x in obj:
            for k in x:
                if k not in cols:
                    cols.append(k)
        tables.append({"name": prefix or "rows", "columns": [str(c) for c in cols], "rows": [[_cell(x.get(c)) for c in cols] for x in obj]})
    elif isinstance(obj, list):
        kv.append((prefix, ", ".join(_cell(x) for x in obj)))
    else:
        kv.append((prefix, _cell(obj)))


def build_document(scope_type: str, items: list) -> dict:
    sections = []
    for it in items:
        kv: list = []
        tables: list = []
        _flatten("", it["data"], kv, tables)
        sections.append({"id": it["id"], "kind": scope_type, "kv": kv, "tables": tables, "data": it["data"]})
    return {"title": f"ParseFusion export ({scope_type}, {len(items)} item(s))", "items": sections}


# ---------------------------------------------------------------------------
# renderers: plugin registry  render(doc, options) -> bytes
# ---------------------------------------------------------------------------
RENDERERS: dict = {}


def renderer(fmt: str, content_type: str, ext: str):
    def deco(fn):
        RENDERERS[fmt] = {"fn": fn, "content_type": content_type, "ext": ext}
        return fn
    return deco


def _csv_safe(s: str) -> str:
    return "'" + s if s[:1] in ("=", "+", "-", "@", "\t", "\r") else s


@renderer("json", "application/json", "json")
def _r_json(doc: dict, opts: dict) -> bytes:
    payload = {"title": doc["title"], "items": [{"id": i["id"], "kind": i["kind"], "data": i["data"]} for i in doc["items"]]}
    return json.dumps(payload, sort_keys=True, ensure_ascii=False, indent=2).encode("utf-8")


@renderer("markdown", "text/markdown; charset=utf-8", "md")
def _r_md(doc: dict, opts: dict) -> bytes:
    esc = lambda s: s.replace("|", "\\|").replace("\n", " ")  # noqa: E731
    out = [f"# {doc['title']}", ""]
    for it in doc["items"]:
        out += [f"## {it['kind']} {it['id']}", ""]
        out += [f"- **{esc(k)}**: {esc(v)}" for k, v in it["kv"]] + [""]
        for t in it["tables"]:
            out += [f"### {t['name']}", "", "| " + " | ".join(esc(c) for c in t["columns"]) + " |", "|" + "---|" * len(t["columns"])]
            out += ["| " + " | ".join(esc(c) for c in r) + " |" for r in t["rows"]] + [""]
    return "\n".join(out).encode("utf-8")


@renderer("csv", "text/csv; charset=utf-8", "csv")
def _r_csv(doc: dict, opts: dict) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\n")
    w.writerow(["item_id", "section", "row", "field", "value"])
    for it in doc["items"]:
        for k, v in it["kv"]:
            w.writerow([it["id"], "fields", "", _csv_safe(k), _csv_safe(v)])
        for t in it["tables"]:
            for i, r in enumerate(t["rows"], 1):
                for c, v in zip(t["columns"], r):
                    w.writerow([it["id"], t["name"], i, _csv_safe(c), _csv_safe(v)])
    return buf.getvalue().encode("utf-8")


@renderer("html", "text/html; charset=utf-8", "html")
def _r_html(doc: dict, opts: dict) -> bytes:
    e = _html.escape
    p = ["<!doctype html><html><head><meta charset='utf-8'><title>", e(doc["title"]), "</title><style>body{font:14px sans-serif;margin:24px}"
         "table{border-collapse:collapse}td,th{border:1px solid #bbb;padding:4px 8px}</style></head><body><h1>", e(doc["title"]), "</h1>"]
    for it in doc["items"]:
        p.append(f"<h2>{e(it['kind'])} {e(it['id'])}</h2><ul>")
        p += [f"<li><b>{e(k)}</b>: {e(v)}</li>" for k, v in it["kv"]]
        p.append("</ul>")
        for t in it["tables"]:
            p.append(f"<h3>{e(t['name'])}</h3><table><tr>" + "".join(f"<th>{e(c)}</th>" for c in t["columns"]) + "</tr>")
            p += ["<tr>" + "".join(f"<td>{e(c)}</td>" for c in r) + "</tr>" for r in t["rows"]]
            p.append("</table>")
    p.append("</body></html>")
    return "".join(p).encode("utf-8")


@renderer("xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "xlsx")
def _r_xlsx(doc: dict, opts: dict) -> bytes:
    try:
        from openpyxl import Workbook
    except ImportError as e:
        raise ec.AgentError("ENGINE_FAILED", "xlsx renderer dependency missing") from e
    wb = Workbook()
    ws = wb.active
    ws.title = "Fields"
    ws.append(["item_id", "field", "value"])
    used = {"Fields"}
    for it in doc["items"]:
        for k, v in it["kv"]:
            ws.append([it["id"], _csv_safe(k), _csv_safe(v)])
        for t in it["tables"]:
            base = re.sub(r"[\[\]\*\?/\\:]", "_", f"{it['id'][:8]}_{t['name']}")[:28] or "table"
            name, n = base, 1
            while name in used:
                n += 1
                name = f"{base[:25]}_{n}"
            used.add(name)
            s = wb.create_sheet(name)
            s.append([_csv_safe(c) for c in t["columns"]])
            for r in t["rows"]:
                s.append([_csv_safe(c) for c in r])
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


@renderer("docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document", "docx")
def _r_docx(doc: dict, opts: dict) -> bytes:
    try:
        from docx import Document
    except ImportError as e:
        raise ec.AgentError("ENGINE_FAILED", "docx renderer dependency missing") from e
    d = Document()
    d.add_heading(doc["title"], 0)
    for it in doc["items"]:
        d.add_heading(f"{it['kind']} {it['id']}", 1)
        for k, v in it["kv"]:
            d.add_paragraph(f"{k}: {v}")
        for t in it["tables"]:
            d.add_heading(t["name"], 2)
            tb = d.add_table(rows=1, cols=max(len(t["columns"]), 1))
            tb.style = "Table Grid"
            for i, c in enumerate(t["columns"]):
                tb.rows[0].cells[i].text = c
            for r in t["rows"]:
                cells = tb.add_row().cells
                for i, v in enumerate(r):
                    cells[i].text = v
    bio = io.BytesIO()
    d.save(bio)
    return bio.getvalue()


@renderer("pdf", "application/pdf", "pdf")
def _r_pdf(doc: dict, opts: dict) -> bytes:
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table as RLTable, TableStyle
    except ImportError as e:
        raise ec.AgentError("ENGINE_FAILED", "pdf renderer dependency missing") from e
    from xml.sax.saxutils import escape as xesc
    st = getSampleStyleSheet()
    bio = io.BytesIO()
    pdf = SimpleDocTemplate(bio, pagesize=landscape(A4), invariant=1, title=doc["title"][:100])
    cl = lambda s: xesc(s if len(s) <= 300 else s[:300] + "...")  # noqa: E731
    story = [Paragraph(xesc(doc["title"]), st["Title"])]
    for it in doc["items"]:
        story.append(Paragraph(xesc(f"{it['kind']} {it['id']}"), st["Heading2"]))
        for k, v in it["kv"]:
            story.append(Paragraph(f"<b>{cl(k)}</b>: {cl(v)}", st["BodyText"]))
        for t in it["tables"]:
            story += [Spacer(1, 6), Paragraph(xesc(t["name"]), st["Heading3"])]
            data = [[Paragraph(cl(c), st["BodyText"]) for c in t["columns"]]] + [[Paragraph(cl(c), st["BodyText"]) for c in r] for r in t["rows"]]
            tb = RLTable(data, repeatRows=1)
            tb.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.4, colors.grey), ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey)]))
            story.append(tb)
    pdf.build(story)
    return bio.getvalue()


# ---------------------------------------------------------------------------
# signed URLs
# ---------------------------------------------------------------------------
def _signed_url(kind: str, export_id: str, user_id: str) -> str:
    exp = int(time.time()) + int(ec.cfg("exports.url_ttl_seconds", 300))
    sig = ec.url_sign(f"{kind}|{export_id}|{exp}|{user_id}")
    return f"/agents/export/{kind}/{export_id}?exp={exp}&uid={quote(user_id, safe='')}&sig={sig}"


def _check_signed(kind: str, export_id: str, exp: int, uid: str, sig: str, user: dict) -> None:
    if exp < int(time.time()):
        raise ec.AgentError("FORBIDDEN", "Download link has expired")
    if not ec.url_verify(f"{kind}|{export_id}|{exp}|{uid}", sig) or uid != user["user_id"]:
        raise ec.AgentError("FORBIDDEN", "Invalid download link")


# ---------------------------------------------------------------------------
# core: create
# ---------------------------------------------------------------------------
def run(inp: ExportIn, user: dict) -> dict:
    t0 = time.perf_counter()
    ec.dbg(AGENT, "start", scope=inp.scope.type, n_ids=len(inp.scope.ids), format=inp.format, user=user["user_id"])
    scopes = ec.cfg("exports.scopes", DEFAULT_SCOPES)
    formats = ec.cfg("exports.formats", DEFAULT_FORMATS)
    if inp.scope.type not in scopes:
        raise ec.AgentError("INVALID_INPUT", "Unsupported scope type", {"allowed": scopes})
    if inp.format not in formats:
        raise ec.AgentError("UNSUPPORTED_FORMAT", "Unsupported export format", {"allowed": formats})
    if inp.format not in RENDERERS:
        raise ec.AgentError("UNSUPPORTED_FORMAT", "No renderer is installed for this format", {"installed": sorted(RENDERERS)})
    ids = list(dict.fromkeys(str(i) for i in inp.scope.ids))
    if not ids:
        raise ec.AgentError("INVALID_INPUT", "scope.ids is empty")
    if len(ids) > int(ec.cfg("exports.max_ids", 500)):
        raise ec.AgentError("TOO_LARGE", "Too many ids in one export")
    opts = inp.options or ExportOptions()
    can_unmask = ec.has_cap(user, "unmask")
    masked = True if (opts.masked is None or not can_unmask) else bool(opts.masked)
    ec.dbg(AGENT, "options", requested_masked=opts.masked, can_unmask=can_unmask, effective_masked=masked)

    kind = (ec.cfg("exports.scope_kinds", {}) or {}).get(inp.scope.type, inp.scope.type)
    items, missing = [], []
    for i in ids:
        obj = ec.store().get(kind, i)
        if not obj or (isinstance(obj, dict) and obj.get("tenant_id") and obj["tenant_id"] != user.get("tenant_id")):
            missing.append(i)
        else:
            items.append({"id": i, "data": obj})
    ec.dbg(AGENT, "scope_loaded", found=len(items), missing=len(missing))
    if missing:
        raise ec.AgentError("NOT_FOUND", "Some items were not found", {"missing": missing[:50]})

    resource = (ec.cfg("exports.scope_resources", {}) or {}).get(inp.scope.type, inp.scope.type)
    try:
        locked = {c.lower() for c in ec.import_agent("23_access_control").locked_field_names(user, resource)}
    except ec.AgentError:
        locked = set()  # data DB not configured -> no registry -> nothing registered as lockable
    ev_keys = {k.lower() for k in ec.cfg("exports.evidence_keys", DEFAULT_EVIDENCE_KEYS)}
    drop = set(locked) | (set() if opts.include_evidence else ev_keys)
    sensitive = {k.lower() for k in ec.cfg("exports.sensitive_keys", DEFAULT_SENSITIVE_KEYS)}
    for it in items:
        d = _walk(it["data"], lambda k: k.lower() in drop)
        it["data"] = _mask(d, sensitive) if masked else d
    ec.dbg(AGENT, "filtered", locked=len(locked), evidence_kept=opts.include_evidence, masked=masked)

    doc = build_document(inp.scope.type, items)
    n_rows = sum(len(t["rows"]) for s in doc["items"] for t in s["tables"]) + sum(len(s["kv"]) for s in doc["items"])
    if n_rows > int(ec.cfg("exports.max_rows", 200000)):
        raise ec.AgentError("TOO_LARGE", "Export is too large for a synchronous run", {"rows": n_rows})
    plug = RENDERERS[inp.format]
    data = plug["fn"](doc, {"masked": masked})
    chash = ec.sha256_hex(data)
    ec.dbg(AGENT, "rendered", bytes=len(data), rows=n_rows, content_hash=chash[:12])

    export_id, created = str(uuid.uuid4()), ec.iso_now()
    effective = {"masked": masked, "masked_forced": bool(opts.masked is False and not can_unmask), "include_evidence": opts.include_evidence}
    manifest = {"export_id": export_id, "format": inp.format, "scope": {"type": inp.scope.type, "ids": ids}, "options": effective,
                "inputs": [{"id": it["id"], "sha256": ec.sha256_hex(ec.canon(it["data"]))} for it in items], "locked_columns_excluded": sorted(locked),
                "agent_version": VERSION, "content_hash": chash, "created_at": created, "created_by": user["user_id"]}
    msig = ec.sign(ec.canon(manifest))
    aad = f"{user.get('tenant_id', 'default')}:{export_id}".encode()
    blob = ec.encrypt(data, aad)
    filename = f"export-{export_id[:8]}.{plug['ext']}"
    ec.store().put("export_blob", export_id, {"b64": base64.b64encode(blob).decode("ascii"), "content_type": plug["content_type"], "filename": filename})
    ec.store().put("export_manifest", export_id, {"manifest": manifest, "signature": msig})
    ec.dbg(AGENT, "stored_encrypted", blob_bytes=len(blob), kid=msig["kid"])

    with ec.app_engine().begin() as c:
        c.execute(insert(exports_tbl).values(export_id=export_id, tenant_id=user.get("tenant_id", "default"), user_id=user["user_id"],
                                             scope=ec.jdumps(manifest["scope"]), format=inp.format, content_hash=chash, size_bytes=len(data),
                                             masked=1 if masked else 0, created_at=created, content_type=plug["content_type"], filename=filename))
        ec.audit("export_created", "export", export_id, "success", {"scope_type": inp.scope.type, "n_ids": len(ids), "format": inp.format,
                 "masked": masked, "content_hash": chash, "bytes": len(data), "kid": msig["kid"]}, strict=True, user=user)
    sensitive_scope = inp.scope.type in ec.cfg("exports.sensitive_scopes", [])
    if (not masked) or sensitive_scope:
        ec.notify_send("export_sensitive", message=f"Export {export_id} by {user['user_id']} (role {user['role']}): scope {inp.scope.type}, "
                       f"{'UNMASKED' if not masked else 'masked'}, format {inp.format}.", link="/admin/exports")
    ec.dbg(AGENT, "done", export_id=export_id, ms=round((time.perf_counter() - t0) * 1000, 1))
    return {"export_id": export_id, "format": inp.format, "download_url": _signed_url("download", export_id, user["user_id"]),
            "content_hash": chash, "signed_manifest_url": _signed_url("manifest", export_id, user["user_id"]), "created_at": created}


def history(user: dict, limit: int = 100) -> dict:
    lim = min(max(int(limit), 1), 500)
    with ec.app_engine().connect() as c:
        rows = c.execute(select(exports_tbl).where(exports_tbl.c.user_id == user["user_id"], exports_tbl.c.tenant_id == user.get("tenant_id", "default"))
                         .order_by(exports_tbl.c.created_at.desc()).limit(lim)).all()
    out = [{"export_id": r.export_id, "format": r.format, "scope": ec.jloads(r.scope, {}), "content_hash": r.content_hash, "size_bytes": r.size_bytes,
            "masked": bool(r.masked), "created_at": r.created_at, "download_url": _signed_url("download", r.export_id, user["user_id"]),
            "signed_manifest_url": _signed_url("manifest", r.export_id, user["user_id"])} for r in rows]
    ec.audit("export_history_viewed", "export", "history", "success", {"returned": len(out)}, user=user)
    ec.dbg(AGENT, "history", returned=len(out))
    return {"exports": out}


def _load_owned(export_id: str, user: dict):
    with ec.app_engine().connect() as c:
        r = c.execute(select(exports_tbl).where(exports_tbl.c.export_id == export_id, exports_tbl.c.tenant_id == user.get("tenant_id", "default"),
                                                exports_tbl.c.user_id == user["user_id"])).first()
    if not r:
        raise ec.AgentError("NOT_FOUND", "Export not found")
    return r


def download(export_id: str, exp: int, uid: str, sig: str, user: dict) -> Response:
    ec.dbg(AGENT, "download_start", export_id=export_id)
    _check_signed("download", export_id, exp, uid, sig, user)
    row = _load_owned(export_id, user)
    stored = ec.store().get("export_blob", export_id)
    if not stored:
        raise ec.AgentError("NOT_FOUND", "Export content is no longer available")
    try:
        data = ec.decrypt(base64.b64decode(stored["b64"]), f"{row.tenant_id}:{export_id}".encode())
    except Exception as e:  # tampered ciphertext / wrong AAD / missing key -> fail closed
        ec.audit("export_downloaded", "export", export_id, "error", {"reason": "decrypt_failed"}, user=user)
        raise ec.AgentError("ENGINE_FAILED", "Stored export could not be decrypted") from e
    if ec.sha256_hex(data) != row.content_hash:
        ec.audit("export_downloaded", "export", export_id, "error", {"reason": "hash_mismatch"}, user=user)
        raise ec.AgentError("ENGINE_FAILED", "Stored export failed its integrity check")
    ec.audit("export_downloaded", "export", export_id, "success", {"bytes": len(data), "content_hash": row.content_hash}, strict=True, user=user)
    ec.dbg(AGENT, "download_ok", bytes=len(data))
    return Response(content=data, media_type=row.content_type, headers={"Content-Disposition": f'attachment; filename="{row.filename}"',
                                                                         "Cache-Control": "no-store", "X-Content-SHA256": row.content_hash})


def manifest(export_id: str, exp: int, uid: str, sig: str, user: dict) -> dict:
    _check_signed("manifest", export_id, exp, uid, sig, user)
    _load_owned(export_id, user)
    m = ec.store().get("export_manifest", export_id)
    if not m:
        raise ec.AgentError("NOT_FOUND", "Manifest not found")
    ec.audit("export_manifest_viewed", "export", export_id, "success", {}, user=user)
    return m


# ---------------------------------------------------------------------------
# routes
# ---------------------------------------------------------------------------
router = APIRouter()


@router.post("/agents/export")
def post_export(request: Request, body: dict = Body(default=None)):
    return ec.handle(request, lambda: run(ExportIn.model_validate(body or {}), ec.current_user()))


@router.get("/agents/export/history")
def get_history(request: Request, limit: int = Query(100)):
    return ec.handle(request, lambda: history(ec.current_user(), limit))


@router.get("/agents/export/download/{export_id}")
def get_download(request: Request, export_id: str, exp: int = Query(...), uid: str = Query(...), sig: str = Query(...)):
    ctx = ec.begin_request(request)
    try:
        return download(export_id, exp, uid, sig, ec.current_user())
    except ec.AgentError as e:
        return ec.error_response(e, ctx.request_id)
    except Exception:
        return ec.error_response(ec.AgentError("ENGINE_FAILED", "Internal error"), ctx.request_id)


@router.get("/agents/export/manifest/{export_id}")
def get_manifest(request: Request, export_id: str, exp: int = Query(...), uid: str = Query(...), sig: str = Query(...)):
    return ec.handle(request, lambda: manifest(export_id, exp, uid, sig, ec.current_user()))
